from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

import hound_core.api as api_module
from hound_core.api import app


def test_resume_parse_endpoint_returns_profile() -> None:
    doc = Document()
    doc.add_paragraph("Backend developer with Python and SQL")

    payload = BytesIO()
    doc.save(payload)

    client = TestClient(app)
    response = client.post(
        "/resume/parse",
        files={
            "file": (
                "resume.docx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert "profile" in body
    assert "python" in body["profile"]["skills"]


def test_resume_parse_endpoint_rejects_non_document_file() -> None:
    client = TestClient(app)
    response = client.post(
        "/resume/parse",
        files={"file": ("resume.txt", b"plain text", "text/plain")},
    )
    assert response.status_code == 400


def test_resume_parse_endpoint_uses_cache_for_same_file(monkeypatch, tmp_path) -> None:
    parse_calls = {"count": 0}

    def fake_parse_resume_file(filename: str, content: bytes):
        parse_calls["count"] += 1
        return {
            "profile": {
                "skills": ["python"],
                "experiences": [],
                "source": filename,
            },
            "extracted_text": "python resume content",
            "used_llm": True,
        }

    monkeypatch.setattr(api_module, "parse_resume_file", fake_parse_resume_file)
    monkeypatch.setenv("HOUND_RESUME_CACHE_PATH", str(tmp_path / "resume-cache.json"))

    doc = Document()
    doc.add_paragraph("Backend developer with Python")
    payload = BytesIO()
    doc.save(payload)
    file_bytes = payload.getvalue()

    client = TestClient(app)

    first = client.post(
        "/resume/parse",
        files={
            "file": (
                "resume.docx",
                file_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    second = client.post(
        "/resume/parse",
        files={
            "file": (
                "resume.docx",
                file_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["cache_hit"] is False
    assert second.json()["cache_hit"] is True
    assert parse_calls["count"] == 1


def test_resume_parse_cache_keeps_multiple_hash_entries(monkeypatch, tmp_path) -> None:
    parse_calls = {"count": 0}

    def fake_parse_resume_file(filename: str, content: bytes):
        parse_calls["count"] += 1
        return {
            "profile": {
                "skills": ["python"],
                "experiences": [],
                "source": filename,
            },
            "extracted_text": "python resume content",
            "used_llm": True,
        }

    monkeypatch.setattr(api_module, "parse_resume_file", fake_parse_resume_file)
    monkeypatch.setenv("HOUND_RESUME_CACHE_PATH", str(tmp_path / "resume-cache.json"))
    monkeypatch.setenv("HOUND_RESUME_CACHE_MAX_ENTRIES", "4")

    doc1 = Document()
    doc1.add_paragraph("Resume one")
    payload1 = BytesIO()
    doc1.save(payload1)

    doc2 = Document()
    doc2.add_paragraph("Resume two")
    payload2 = BytesIO()
    doc2.save(payload2)

    client = TestClient(app)

    first = client.post(
        "/resume/parse",
        files={"file": ("resume-1.docx", payload1.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    second = client.post(
        "/resume/parse",
        files={"file": ("resume-2.docx", payload2.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    third = client.post(
        "/resume/parse",
        files={"file": ("resume-1.docx", payload1.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert first.status_code == 200
    assert second.status_code == 200
    assert third.status_code == 200
    assert first.json()["cache_hit"] is False
    assert second.json()["cache_hit"] is False
    assert third.json()["cache_hit"] is True
    assert parse_calls["count"] == 2
